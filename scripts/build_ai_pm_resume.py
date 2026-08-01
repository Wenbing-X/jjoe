from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "documents" / "谢文炳_AI项目经理_优化简历.docx"

FONT = "Microsoft YaHei"
INK = RGBColor(28, 32, 40)
MUTED = RGBColor(102, 109, 122)
ACCENT = RGBColor(64, 105, 238)
SOFT_ACCENT = "EEF3FF"


def set_run(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic
    return run


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Resume Section")
    p.add_run(text)
    set_keep(p, keep_next=True)
    return p


def add_bullet(doc, prefix, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Mm(5.5)
    p.paragraph_format.first_line_indent = Mm(-3.0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.08
    set_run(p.add_run(prefix), size=9.1, bold=True, color=INK)
    set_run(p.add_run(text), size=9.1, color=INK)
    set_keep(p)
    return p


def add_role_line(doc, title, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    set_run(p.add_run(title), size=10.3, bold=True, color=INK)
    set_run(p.add_run("  |  "), size=9.3, color=MUTED)
    set_run(p.add_run(meta), size=9.0, color=MUTED)
    set_keep(p, keep_next=True)
    return p


def build_resume():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(12.5)
    section.bottom_margin = Mm(12.5)
    section.left_margin = Mm(14.5)
    section.right_margin = Mm(14.5)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(7)

    # compact_reference_guide preset with named A4 resume overrides.
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.1

    list_style = doc.styles["List Bullet"]
    list_style.font.name = FONT
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    list_style.font.size = Pt(9.1)
    list_style.font.color.rgb = INK

    section_style = doc.styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    section_style.font.name = FONT
    section_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    section_style.font.size = Pt(11.2)
    section_style.font.bold = True
    section_style.font.color.rgb = ACCENT
    section_style.paragraph_format.space_before = Pt(7)
    section_style.paragraph_format.space_after = Pt(4)
    section_style.paragraph_format.line_spacing = 1.0

    doc.core_properties.title = "谢文炳 - AI项目经理求职简历"
    doc.core_properties.subject = "AI项目经理 / AI内容项目经理"
    doc.core_properties.author = "谢文炳"
    doc.core_properties.keywords = "AI项目经理, AIGC, ComfyUI, 内容项目, 项目管理"

    name = doc.add_paragraph()
    name.paragraph_format.space_before = Pt(0)
    name.paragraph_format.space_after = Pt(1)
    set_run(name.add_run("谢文炳"), size=25, bold=True, color=INK)
    set_keep(name, keep_next=True)

    target = doc.add_paragraph()
    target.paragraph_format.space_before = Pt(0)
    target.paragraph_format.space_after = Pt(3)
    set_run(target.add_run("AI 项目经理 / AI 内容项目经理（初级）"), size=11, bold=True, color=ACCENT)
    set_keep(target, keep_next=True)

    contact = doc.add_paragraph()
    contact.paragraph_format.space_before = Pt(0)
    contact.paragraph_format.space_after = Pt(7)
    contact.paragraph_format.line_spacing = 1.0
    set_run(
        contact.add_run("长沙  |  +86 181 5546 6798  |  wuzibx@foxmail.com"),
        size=9.1,
        color=MUTED,
    )

    add_section_heading(doc, "职业概述")
    summary = doc.add_paragraph()
    summary.paragraph_format.left_indent = Mm(3.2)
    summary.paragraph_format.right_indent = Mm(3.2)
    summary.paragraph_format.space_before = Pt(2.5)
    summary.paragraph_format.space_after = Pt(3.5)
    summary.paragraph_format.line_spacing = 1.18
    shade_paragraph(summary, SOFT_ACCENT)
    set_run(
        summary.add_run(
            "工商管理本科，具备 AI 内容项目从 0 到 1 的独立执行经验。围绕 AI 短剧账号完成定位与选题、脚本拆解、生成式制作、发布运营和数据复盘，能够将模糊目标拆解为任务与交付物，并协调工具与内容流程推进交付。熟悉 ComfyUI、即梦、剪映、Photoshop 及抖音 / 小红书运营，希望从 AI 内容项目和 AI 应用项目切入项目管理岗位。"
        ),
        size=9.2,
        color=INK,
    )
    set_keep(summary)

    add_section_heading(doc, "AI 项目实践")
    add_role_line(doc, "AI 短剧账号 0-1 搭建与运营", "个人项目")
    add_bullet(
        doc,
        "项目规划：",
        "围绕账号定位、目标平台与内容方向拆解制作链路，建立“选题 - 脚本 - 出图 - 视频 - 剪辑 - 发布 - 复盘”的端到端流程。",
    )
    add_bullet(
        doc,
        "AI 工作流：",
        "本地部署 ComfyUI，通过节点搭建文生图、图生视频工作流，并结合即梦等 AI 视频工具完成素材生产，提升内容制作效率。",
    )
    add_bullet(
        doc,
        "交付推进：",
        "制定内容排期与素材清单，按节点推进脚本、画面、剪辑和发布，沉淀可复用的制作流程与模板。",
    )
    add_bullet(
        doc,
        "数据迭代：",
        "围绕播放量、互动率等指标复盘选题、叙事节奏和视觉表现，持续优化内容策略。",
    )

    add_section_heading(doc, "校园与组织经历")
    add_role_line(doc, "学生会 · 大二实习干事", "2024.01 - 2025.07")
    add_bullet(
        doc,
        "活动执行：",
        "参与校园活动策划与落地，协助拆解任务、准备物料、跟进现场流程与协同事项，保障活动按计划完成。",
    )

    add_section_heading(doc, "教育背景与竞赛")
    add_role_line(doc, "湖南涉外经济学院 · 工商管理 · 本科", "预计 2026.06 毕业")
    add_bullet(
        doc,
        "竞赛经历：",
        "ERP 沙盘模拟大赛（2025.06），参与经营决策、资源配置与团队协作模拟。",
    )

    add_section_heading(doc, "专业技能")
    add_bullet(
        doc,
        "项目管理：",
        "需求拆解、WBS 任务排期、里程碑跟踪、风险 / 问题清单、项目复盘；可编写基础 PRD、项目计划和周报。",
    )
    add_bullet(
        doc,
        "AI 与内容：",
        "ComfyUI 本地部署与节点工作流、即梦、Prompt 设计、文生图 / 图生视频、AI 短剧脚本与全流程制作。",
    )
    add_bullet(
        doc,
        "工具与平台：",
        "剪映、Photoshop、WPS / Office（Word、Excel、PPT）、飞书文档与多维表格基础；熟悉抖音、小红书内容运营。",
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    set_run(footer.add_run("谢文炳 · AI 项目经理求职简历 · 2026"), size=7.5, color=MUTED)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_resume()
